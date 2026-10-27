"""Adapter between the local SEB (Malone-Breneman) shear wall spreadsheet and the engine.

The spreadsheet itself is private and git-ignored. This module holds only cell
addresses; every number is read from the local copy at run time. Set
``SEB_XLSX`` to its path, or keep it at the default location under references/.

Sheet layout (both supported sheets): per-level tables list level 5 first (top)
and level 1 last (bottom); spreadsheet level i is engine story i.
"""

from __future__ import annotations

import csv
import dataclasses
import os
import re
import shutil
from dataclasses import dataclass
from fractions import Fraction
from pathlib import Path

from design.components.schema import RECORD_TYPES
from design.inputs import ChordSpec, DesignInputs, RodSystem, StoryInput

REPO = Path(__file__).resolve().parents[1]
DEFAULT_PATH = REPO / "references" / "SEB Master Shear Wall Design-2025-8 5 over 1 Match V2.xlsb.xlsx"
EXAMPLE_COMPONENTS = Path(__file__).parent / "fixtures" / "woodworks_components"

SHEETS = ("5 Story 2025 Example Wall", "SDPWS 5 story example")
N = 5  # levels per sheet
SOURCE = "SEB shear wall spreadsheet (local copy, not in Git)"

# Deformation components the engine models, in the sheet's "In da" switch row 214 (F..M):
# slack/take-up, shrinkage, rod, bearing plate, tension-side top of post, bottom plate, sheathing, top plate.
ENGINE_DA_SWITCHES = (1, 0, 1, 1, 0, 1, 1, 1)


def workbook_path() -> Path | None:
    p = Path(os.environ.get("SEB_XLSX", DEFAULT_PATH))
    return p if p.exists() else None


def load(path: Path):
    import openpyxl
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl warns about the workbook's data validation extensions
        return openpyxl.load_workbook(path, data_only=True)


def row(top_row: int, level: int) -> int:
    """Row of ``level`` in a per-level table whose level-5 row is ``top_row``."""
    return top_row + (N - level)


def num(text) -> float:
    """Parse '3', '3.5', '5/8', '3 3/8', '1 5/16' into a float."""
    return float(sum(Fraction(part) for part in str(text).split()))


# --------------------------------------------------------------------------
# Design Tables lookups
# --------------------------------------------------------------------------

def _table(wb, first: int, last: int) -> dict[str, int]:
    dt = wb["Design Tables"]
    return {str(dt[f"B{r}"].value): r for r in range(first, last + 1) if dt[f"B{r}"].value is not None}


@dataclass(frozen=True)
class Plate:
    width_in: float
    length_in: float
    thickness_in: float | None
    hole_in: float | None  # wood bearing hole, from the catalog name


def parse_plate(name: str) -> Plate:
    """'3x4x5/8-1 5/16' -> W 3, L 4, t 5/8, hole 1 5/16; thickness and hole may be absent."""
    dims_text, _, hole_text = name.partition("-")
    dims = dims_text.split("x")
    w, l = num(dims[0]), num(dims[1])
    t = num(dims[2]) if len(dims) > 2 else None
    return Plate(w, l, t, num(hole_text) if hole_text.strip() else None)


def parse_spacing(nail_text: str) -> float:
    return num(re.sub(r'^\(\d\)\s*|"', "", str(nail_text)).strip())


def component_dir(wb, dest: Path) -> Path:
    """Write a component library built from the sheet's Design Tables into ``dest``.

    Lumber, nails and panels come from the WoodWorks fixture (the sheet does not
    tabulate F_c or E_min). Nominal shears are SDPWS 2015 values in the sheet with
    an ASD divisor of 2.0; they are rescaled to v_n = 2.8 (v_s / 2) so the engine's
    Ω = 2.8 reproduces the sheet's allowable. Double-sided types are stored per side.
    """
    dest.mkdir(parents=True, exist_ok=True)
    for f in EXAMPLE_COMPONENTS.glob("*.csv"):
        shutil.copy(f, dest / f.name)
    dt = wb["Design Tables"]
    common = dict(source_doc=SOURCE, page="", status="memory", entered_by="test", checked_by="")

    asm = []
    for t, r in _table(wb, 7, 34).items():
        double = str(dt[f"G{r}"].value).startswith("(2)")
        per_side = 2 if double else 1
        vs, ga = float(dt[f"I{r}"].value) / per_side, float(dt[f"J{r}"].value) / per_side
        ply = t.endswith("P")
        asm.append(dict(
            assembly_id=f"SEB-{t}", sheathing_material="WSP", panel_grade=str(dt[f"D{r}"].value),
            min_nominal_panel_thickness_label="15/32", min_nominal_panel_thickness_in=0.46875,
            min_fastener_penetration_in="", nail_designation="10d common",
            edge_nail_spacing_in=parse_spacing(dt[f"G{r}"].value),
            v_seismic_nominal_plf=2.8 * vs / 2, Ga_osb_kips_per_in="" if ply else ga,
            Ga_plywood_kips_per_in=ga if ply else "", v_wind_nominal_plf="", footnotes="",
            table="Design Tables B7:M34", notes="v_n rescaled from SDPWS 2015 v_s: 2.8 (v_s / 2)", **common))
    rods = []
    for t, r in _table(wb, 80, 97).items():
        dia = str(dt[f"C{r}"].value)
        rods.append(dict(
            system_id=f"SEB-ROD-{t}", manufacturer="per spreadsheet", rod_diameter_in=num(dia.rstrip("H")),
            rod_material="high strength" if dia.endswith("H") else "standard",
            rod_allowable_tension_lb=1000 * float(dt[f"D{r}"].value), rod_net_area_in2=float(dt[f"F{r}"].value),
            take_up_device="", bearing_plate="", evaluation_report="", table="Design Tables B80:G97",
            notes="", **common))
    for name, rows in (("shear_wall_assemblies", asm), ("continuous_rod_systems", rods)):
        fields = [f.name for f in dataclasses.fields(RECORD_TYPES[name])]
        with (dest / f"{name}.csv").open("w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=fields)
            w.writeheader()
            w.writerows(rows)
    return dest


# --------------------------------------------------------------------------
# Inputs
# --------------------------------------------------------------------------

def unsupported(wb, sheet: str) -> str | None:
    """Reason the sheet's settings fall outside what the engine models, or None."""
    ws = wb[sheet]
    if ws["H204"].value != "Ashley":
        return "H204 is not 'Ashley' (the WoodWorks-matching load path)"
    da = tuple(int(ws.cell(214, c).value or 0) for c in range(6, 14))
    if da != ENGINE_DA_SWITCHES:
        return f"'In da' switches F214:M214 = {da}, engine models {ENGINE_DA_SWITCHES}"
    if any(ws.cell(215, c).value for c in range(6, 14)):
        return "rotation carried up the wall (row 215) is on; engine does not model it (D6)"
    if ws["J33"].value not in (1, 1.0):
        return "redundancy factor ρ ≠ 1; engine has no ρ input"
    posts = {str(ws[f"D{row(62, i)}"].value) for i in range(1, N + 1)}
    if not all(p.startswith("N") for p in posts):
        return f"post types {posts}: only the 2x4 'N' types are mapped"
    return None


def inputs_from_sheet(wb, sheet: str, run: str = "design") -> DesignInputs:
    """Engine inputs for one sheet.

    ``run="design"`` uses the strength-design posts and rods (D62, H85);
    ``run="drift"`` uses the deflection table's own selections (I112, J112),
    which a sheet may set differently.
    """
    ws, dt = wb[sheet], wb["Design Tables"]
    post_col, post_top, rod_col, rod_top = ("D", 62, "H", 85) if run == "design" else ("I", 112, "J", 112)
    posts, plates = _table(wb, 54, 74), _table(wb, 105, 120)
    stories = []
    for i in range(1, N + 1):
        above = i + 1 if i < N else None
        wall = float(ws[f"C{row(50, i)}"].value)
        at_level = float(ws[f"D{row(50, i)}"].value)
        LL, Lr = float(ws[f"F{row(50, i)}"].value), float(ws[f"G{row(50, i)}"].value)
        LL_up = float(ws[f"F{row(50, above)}"].value) if above else 0.0
        Lr_up = float(ws[f"G{row(50, above)}"].value) if above else 0.0

        nail_type = str(ws[f"H{row(39, i)}"].value)
        double = str(dt[f"G{_table(wb, 7, 34)[nail_type]}"].value).startswith("(2)")
        p = posts[str(ws[f"{post_col}{row(post_top, i)}"].value)]
        plate = parse_plate(str(ws[f"D{row(97, i)}"].value))
        slack = float(ws[f"C{row(209, i)}"].value)
        takeup = float(ws[f"E{row(209, i)}"].value)
        stories.append(StoryInput(
            h_ft=float(ws[f"J{row(16, i)}"].value),
            F_lb=1000 * float(ws[f"J{row(25, i)}"].value),
            assembly_id=f"SEB-{nail_type}", sides=2 if double else 1,
            chord=ChordSpec(n_posts=round(float(dt[f"G{p}"].value) / 5.25),
                            chord_length_ft=float(dt[f"F{p}"].value) / 12,
                            d_arm_ft=float(ws[f"I{row(50, i)}"].value)),
            tiedown=RodSystem(f"SEB-ROD-{ws[f'{rod_col}{row(rod_top, i)}'].value}", plate.width_in,
                              plate.length_in, plate.thickness_in or 0.375, slack_in=slack,
                              takeup_elongation_in=takeup, bearing_hole_in=plate.hole_in),
            w_D_plf=at_level - wall, w_wall_plf=wall, w_L_plf=LL - LL_up, w_Lr_plf=Lr - Lr_up,
        ))
    return DesignInputs(
        load_type="seismic", b_ft=float(ws["E32"].value), stories=tuple(stories),
        S_DS=float(ws["J32"].value), C_d=4.0, I_e=1.0,  # R222 = O222*4; no I_e term in the sheet
        risk_category="II", sdc="D",  # S222 = 0.02 h_sx
        wood_bearing_factor=float(ws["D175"].value), project_name=f"SEB spreadsheet: {sheet}",
    )


# --------------------------------------------------------------------------
# Expected values
# --------------------------------------------------------------------------

@dataclass(frozen=True)
class Pair:
    """One engine quantity and the spreadsheet cell it is compared with."""

    label: str
    kind: str  # "step" or "check"
    engine_id: str
    attr: str  # for checks: demand | capacity | dcr
    cell: str
    scale: float = 1.0  # engine value = scale * cell value
    known_difference: str = ""  # reason, when a mismatch is expected
    run: str = "design"  # which engine run supplies the value (see inputs_from_sheet)


def pairs(i: int) -> list[Pair]:
    """Engine vs spreadsheet quantities for story ``i``."""
    r = lambda top: row(top, i)  # noqa: E731
    S = f"S{i}"
    return [
        Pair("unit shear v (ASD)", "step", f"{S}.shear.v", "", f"D{r(39)}"),
        Pair("allowable shear", "step", f"{S}.shear.v_allow", "", f"I{r(39)}"),
        Pair("shear DCR", "check", f"{S}.shear", "dcr", f"J{r(39)}"),
        Pair("dead load to chord P_D", "step", f"{S}.chord.P_D", "", f"J{r(50)}", 1000),
        Pair("T-C couple C_E", "step", f"{S}.chord.C_E", "", f"M{r(50)}", 1000),
        Pair("compression demand", "check", f"{S}.comp", "demand", f"C{r(62)}", 1000),
        Pair("compression capacity", "check", f"{S}.comp", "capacity", f"E{r(62)}", 1000),
        Pair("compression DCR", "check", f"{S}.comp", "dcr", f"F{r(62)}"),
        Pair("overturning M_OT", "step", f"{S}.statics.M_OT", "", f"E{r(85)}", 1000),
        Pair("resisting moment M_R", "step", f"{S}.chord.M_R", "", f"C{r(85)}", -1000),
        Pair("net moment (ASD)", "step", f"{S}.chord.asd.net", "", f"F{r(85)}", -1000),
        Pair("chord tension T (ASD)", "step", f"{S}.chord.asd.T", "", f"G{r(85)}", 1000),
        Pair("rod capacity", "check", f"{S}.rod", "capacity", f"I{r(85)}", 1000),
        Pair("rod DCR", "check", f"{S}.rod", "dcr", f"J{r(85)}"),
        Pair("differential tension ΔT", "step", f"{S}.chord.dT", "", f"K{r(85)}", 1000),
        Pair("bearing area A_brg", "step", f"{S}.plate_brg.A_brg", "", f"E{r(97)}", known_difference=PLATE_TABLE),
        Pair("bearing factor C_b", "step", f"{S}.plate_brg.C_b", "", f"F{r(97)}"),
        Pair("plate bearing capacity", "check", f"{S}.plate_brg", "capacity", f"G{r(97)}", 1000),
        Pair("plate bearing DCR", "check", f"{S}.plate_brg", "dcr", f"H{r(97)}"),
        *drift_pairs(i),
    ]


PLATE_TABLE = "plate areas in the sheet's Design Tables are flagged 'Areas not verified' (D103)"


def drift_pairs(i: int) -> list[Pair]:
    r = lambda top: row(top, i)  # noqa: E731
    S = f"S{i}"
    pairs_ = [
        Pair("drift unit shear v", "step", f"{S}.drift.v", "", f"E{r(112)}"),
        Pair("drift chord tension", "step", f"{S}.chord.drift.T", "", f"J{r(197)}", -1000),
        Pair("drift ΔT", "step", f"{S}.chord.dT_drift", "", f"L{r(197)}", -1000),
        Pair("crushing: bottom plate", "step", f"{S}.drift.crush_bot", "", f"K{r(209)}"),
        Pair("crushing: sheathing", "step", f"{S}.drift.crush_sh", "", f"L{r(209)}"),
        Pair("crushing: top plate", "step", f"{S}.drift.crush_top", "", f"M{r(209)}"),
        Pair("rod elongation", "step", f"{S}.drift.rod", "", f"H{r(209)}"),
        Pair("bearing plate crushing", "step", f"{S}.drift.crush_plate", "", f"I{r(209)}"),
        Pair("Δ_a at wall ends", "step", f"{S}.drift.delta_a", "", f"O{r(209)}"),
        Pair("bending term", "step", f"{S}.drift.d1", "", f"R{r(122)}"),
        Pair("shear/slip term", "step", f"{S}.drift.d2", "", f"C{r(222)}"),
        Pair("anchorage term", "step", f"{S}.drift.d3", "", f"P{r(209)}"),
        Pair("δ_sw", "step", f"{S}.drift.delta_sw", "", f"O{r(222)}",
             known_difference="sheet totals use the post-inertia bending term (H122) instead of SDPWS 8vh³/EAb"),
        Pair("design story drift", "check", f"{S}.drift", "demand", f"R{r(222)}",
             known_difference="C_d × δ_sw; follows the δ_sw bending-term difference"),
        Pair("allowable drift", "check", f"{S}.drift", "capacity", f"S{r(222)}"),
    ]
    return [dataclasses.replace(p, run="drift") for p in pairs_]


# Quantities downstream of the bearing-plate area; differ where the sheet's plate table area is off.
PLATE_CHAIN = {"plate bearing capacity", "plate bearing DCR", "bearing plate crushing", "Δ_a at wall ends",
               "anchorage term", "δ_sw", "design story drift"}


def known_reason(wb, sheet: str, i: int, p: Pair) -> str:
    """Why a mismatch on ``p`` is expected, or "" if it is not."""
    if p.known_difference:
        return p.known_difference
    if p.label in PLATE_CHAIN:
        plate = str(wb[sheet][f"D{row(97, i)}"].value)
        dt = wb["Design Tables"]
        r = _table(wb, 105, 120).get(plate)
        if r is not None:
            g = parse_plate(plate)
            hole = g.hole_in or 0.0
            computed = g.width_in * g.length_in - 3.141592653589793 * hole**2 / 4
            if abs(computed - float(dt[f"E{r}"].value)) > 0.01 * computed:
                return f"downstream of plate '{plate}': {PLATE_TABLE}"
    return ""


def engine_value(result, p: Pair):
    """The engine's number for ``p``, or None when the step or check did not produce one."""
    try:
        if p.kind == "step":
            return result.step(p.engine_id).value
        return getattr(result.check(p.engine_id), p.attr)
    except KeyError:
        return None
